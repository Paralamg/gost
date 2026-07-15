"""Вывод сетки таблицы в документ Word."""

from docx.document import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.shared import Cm, Emu, Length, Pt
from docx.table import Table as DocxTable, _Row

from gost.elements.table_grid import Grid

TABLE_STYLE = "Table Grid"
CELL_STYLE = "Table Text"
CAPTION_SIZE = Pt(14)
ROW_NUMBER_WIDTH = Cm(1.2)


def write_caption(
        document: Document,
        text: str,
        *,
        align: WD_ALIGN_PARAGRAPH,
        page_break_before: bool = False,
) -> None:
    """Пишет подпись таблицы отдельным абзацем."""
    paragraph = document.add_paragraph()
    paragraph.alignment = align

    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Cm(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.space_after = Pt(6)
    fmt.keep_with_next = True
    if page_break_before:
        fmt.page_break_before = True

    paragraph.add_run(text).font.size = CAPTION_SIZE


def write_part(
        document: Document,
        grid: Grid,
        rows: list[list[str]],
        *,
        repeat_head: bool = False,
) -> None:
    """Пишет непрерывный блок строк как отдельную таблицу.

    Args:
        grid: Сетка целиком — нужна для ширины и раскладки столбцов.
        rows: Строки этой части, включая её шапку.
        repeat_head: Помечать ли шапку как повторяемую на каждой странице
            («w:tblHeader»). Только для неразрывной таблицы: при ручном
            разбиении шапка пишется в каждую часть явно.
    """
    table = document.add_table(rows=len(rows), cols=grid.width)
    table.style = TABLE_STYLE
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    widths = _column_widths(document, grid)
    _fix_layout(table, widths)

    # Идентификатор стиля резолвится один раз на таблицу: paragraph.style = "Table Text"
    # ищет стиль по имени и перебирает все стили документа на каждой ячейке — на
    # таблице в сотню строк это 85% времени сборки.
    style_id = document.styles[CELL_STYLE].style_id

    for row, values in zip(table.rows, rows):
        _forbid_split(row)
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = width  # Word ориентируется на ширину ячейки (tcW), а не на gridCol
            paragraph = cell.paragraphs[0]
            paragraph._p.get_or_add_pPr().style = style_id
            paragraph.text = value

    if repeat_head:
        for row in table.rows[:len(grid.head_rows)]:
            _repeat_as_header(row)


def _column_widths(document: Document, grid: Grid) -> list[Length]:
    section = document.sections[0]
    total = section.page_width - section.left_margin - section.right_margin

    if not grid.has_row_numbers:
        return [Emu(total // grid.width)] * grid.width

    data_columns = grid.width - 1
    share = Emu((total - ROW_NUMBER_WIDTH) // data_columns)
    return [ROW_NUMBER_WIDTH, *[share] * data_columns]


def _fix_layout(table: DocxTable, widths: list[Length]) -> None:
    """Фиксирует ширины столбцов.

    Без этого Word считает autofit для каждой таблицы отдельно, и части
    разорванной таблицы разъезжаются по ширине. Ширины ячеек проставляет
    write_part — Word смотрит на них, а не на gridCol.
    """
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width


def _forbid_split(row: _Row) -> None:
    """«w:cantSplit» — строка не разрывается между страницами."""
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _repeat_as_header(row: _Row) -> None:
    """«w:tblHeader» — строка повторяется на каждой странице.

    Word повторяет только непрерывный блок строк от начала таблицы.
    """
    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
