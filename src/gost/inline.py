"""Инлайн-разметка: markdown-стили и неразрывный пробел в run'ы Word.

Единая точка разбора «сырого» текста в оформленные фрагменты. Применяется к телу
абзаца и к подписям таблиц/изображений. Список поддерживаемых стилей — INLINE_STYLES;
добавить стиль = добавить в него кортеж.

Поддерживается:
    ``***текст***``   — жирный курсив
    ``**текст**``     — жирный
    ``*текст*``       — курсив
    ``_текст_``       — курсив
    ``__текст__``     — подчёркнутый
    ``~``             — неразрывный пробел (\\xa0)

Ограничения:
    - Экранирования нет: любая тильда в тексте становится неразрывным пробелом.
    - Одиночный маркер без пары остаётся литералом, но текст с двумя такими
      символами (``a * b * c``) может быть ошибочно распознан как курсив —
      обычное ограничение markdown.
"""

import re
from collections.abc import Callable

from docx.text.paragraph import Paragraph
from docx.text.run import Run

NBSP = "\xa0"
NBSP_MARKER = "~"


def _bold(run: Run) -> None:
    run.bold = True


def _italic(run: Run) -> None:
    run.italic = True


def _bold_italic(run: Run) -> None:
    run.bold = True
    run.italic = True


def _underline(run: Run) -> None:
    run.underline = True


# (открывающий маркер, закрывающий маркер, как оформить run).
# Порядок важен: более длинные/специфичные маркеры идут первыми, иначе «*»
# перехватит содержимое «**»/«***». Альтернация в regex упорядочена, поэтому
# первый подходящий вариант выигрывает.
INLINE_STYLES: list[tuple[str, str, Callable[[Run], None]]] = [
    ("***", "***", _bold_italic),
    ("**", "**", _bold),
    ("*", "*", _italic),
    ("__", "__", _underline),
    ("_", "_", _italic),
]


def _build(
        styles: list[tuple[str, str, Callable[[Run], None]]],
) -> tuple[re.Pattern[str], dict[str, Callable[[Run], None]]]:
    parts, apply_by_name = [], {}
    for i, (open_, close_, apply) in enumerate(styles):
        name = f"s{i}"
        # Уникальная именованная группа на стиль — по match.lastgroup узнаём,
        # какой стиль сработал, и берём внутренний текст.
        parts.append(f"{re.escape(open_)}(?P<{name}>.+?){re.escape(close_)}")
        apply_by_name[name] = apply
    return re.compile("|".join(parts)), apply_by_name


_REGEX, _APPLY = _build(INLINE_STYLES)


def render_inline(paragraph: Paragraph, text: str) -> list[Run]:
    """Разбирает markdown-разметку и `~` (неразрывный пробел) в run'ы Word.

    Возвращает созданные run'ы, чтобы вызывающий код мог доработать их
    (например, задать размер шрифта подписи).
    """
    text = text.replace(NBSP_MARKER, NBSP)
    runs, pos = [], 0
    for m in _REGEX.finditer(text):
        if m.start() > pos:
            runs.append(paragraph.add_run(text[pos:m.start()]))
        name = m.lastgroup
        run = paragraph.add_run(m.group(name))
        _APPLY[name](run)
        runs.append(run)
        pos = m.end()
    if pos < len(text):
        runs.append(paragraph.add_run(text[pos:]))
    return runs
