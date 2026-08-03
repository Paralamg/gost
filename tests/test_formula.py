"""Элемент Formula: вёрстка по ГОСТ 7.32-2017 (6.11)."""

import docx
import pytest
from docx.oxml.ns import qn

from gost.element_factory import ElementFactory
from gost.elements.formula.formula import SEPARATOR, Formula
from gost.index.index_manager import IndexManager
from gost.index.index_type import IndexType
from gost.styles import StyleSheet, apply_gost_styles


@pytest.fixture
def document():
    document = docx.Document()
    apply_gost_styles(document)
    return document


@pytest.fixture
def factory():
    return ElementFactory(IndexManager())


def make(index: str = "1", latex: str = r"E = \frac{mv^2}{2}", **kwargs) -> Formula:
    sheet = StyleSheet.gost()
    return Formula(
        index, latex,
        formula_style=sheet.formula, note_style=sheet.formula_note, **kwargs
    )


def math_nodes(document):
    return document.element.body.findall(f".//{qn('m:oMath')}")


def test_renders_equation_object(document):
    """Формула должна лечь объектом «Уравнение», а не текстом."""
    make().render(document)
    assert len(math_nodes(document)) == 1


def test_math_is_display_equation(document):
    """m:oMathPara — сосед w:r внутри w:p, а не его потомок.

    Обёртка включает выключной режим: без неё Word считает формулу строчной,
    мельчит знак суммы и уводит пределы вбок.
    """
    make().render(document)
    paragraph = document.paragraphs[0]
    assert paragraph._p.find(qn("m:oMathPara")) is not None


def test_paragraph_holds_nothing_but_the_formula(document):
    """Любой обычный run рядом заставил бы Word сжать формулу до строчной."""
    make(index="1").render(document)
    assert document.paragraphs[0].runs == []


def test_numbered_formula_puts_number_inside_equation(document):
    """Номер живёт внутри формулы, разогнанный к правому краю через m:eqArr.

    Так же кодирует нумерованную формулу сам Word при вводе «#(N)».
    """
    make(index="3").render(document)
    math = math_nodes(document)[0]
    row = math.find(f"{qn('m:eqArr')}/{qn('m:e')}")
    assert row is not None
    assert math.find(
        f"{qn('m:eqArr')}/{qn('m:eqArrPr')}/{qn('m:maxDist')}"
    ).get(qn("m:val")) == "1"
    # Номер — объект m:d, круглые скобки у него свои.
    assert "".join(row.find(qn("m:d")).itertext()) == "3"
    assert SEPARATOR in "".join(node.text or "" for node in row.iter(qn("m:t")))


def test_unnumbered_formula_has_no_number(document):
    make(index="").render(document)
    math = math_nodes(document)[0]
    assert math.find(qn("m:eqArr")) is None
    assert SEPARATOR not in "".join(math.itertext())


def test_where_block_format(document):
    """Обозначения встают в колонку: «где» отбивается табом, следующие строки
    начинаются с него же и попадают на ту же позицию табуляции."""
    make(index="1", where={"m": "масса тела, кг", "v": "скорость, м/с"}).render(document)
    notes = [p.text for p in document.paragraphs[1:]]
    assert notes == ["где \tm – масса тела, кг;", "\tv – скорость, м/с."]


def test_single_where_entry_ends_with_period(document):
    make(where={"m": "масса, кг"}).render(document)
    assert document.paragraphs[1].text == "где \tm – масса, кг."


def test_where_keeps_explanation_with_formula(document):
    """Расшифровку нельзя отрывать от формулы переносом страницы."""
    make(where={"m": "масса, кг"}).render(document)
    formula_format = document.paragraphs[0].paragraph_format
    assert formula_format.keep_with_next is True
    assert formula_format.space_after == 0


def test_without_where_keeps_bottom_spacing(document):
    """Без расшифровки ГОСТ требует свободную строку под формулой."""
    make().render(document)
    assert document.paragraphs[0].paragraph_format.space_after > 0


def test_where_supports_inline_markup(document):
    make(where={"*m*": "масса, кг"}).render(document)
    assert any(run.italic for run in document.paragraphs[1].runs)


def test_render_is_repeatable():
    """Автоподбор разрывов рендерит один элемент в несколько документов."""
    formula = make()
    for _ in range(2):
        document = docx.Document()
        apply_gost_styles(document)
        formula.render(document)
        assert len(math_nodes(document)) == 1


def test_broken_latex_fails_on_creation():
    """Ошибка в формуле должна всплыть при создании, а не при сохранении."""
    with pytest.raises(ValueError):
        make(latex=r"\frac{")


def test_factory_numbers_formulas(factory):
    assert [factory.create_formula(r"x").index for _ in range(3)] == ["1", "2", "3"]


def test_factory_skips_counter_for_unnumbered(factory):
    factory.create_formula(r"x")
    assert factory.create_formula(r"y", numbered=False).index == ""
    assert factory.create_formula(r"z").index == "2"


def test_factory_copies_styles(factory):
    formula = factory.create_formula(r"x")
    assert formula.formula_style is not factory.style.formula
    assert formula.note_style is not factory.style.formula_note


def test_chapter_relative_numbering():
    factory = ElementFactory(IndexManager(index_type=IndexType.CHAPTER_RELATIVE))
    factory.create_head("Раздел", True, 1)
    assert factory.create_formula(r"x").index == "1.1"
    factory.create_head("Второй", True, 1)
    assert factory.create_formula(r"y").index == "2.1"


def test_appendix_prefix_numbering():
    factory = ElementFactory(IndexManager(character="А"))
    assert factory.create_formula(r"x").index == "А.1"


def test_formulas_survive_save_and_reopen(tmp_path, factory):
    """Файл должен открываться, а формулы — оставаться на месте."""
    from gost import WordBuilder

    builder = WordBuilder()
    builder.add_element(factory.create_formula(
        r"\sigma = \sqrt{\frac{\sum_{i=1}^{n}(x_i - \bar{x})^2}{n - 1}}",
        where={"n": "объём выборки"},
    ))
    builder.add_element(factory.create_formula(r"a^2 + b^2 = c^2", numbered=False))
    path = tmp_path / "formulas.docx"
    builder.save(path)

    assert len(math_nodes(docx.Document(str(path)))) == 2
